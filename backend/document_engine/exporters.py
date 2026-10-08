from io import BytesIO
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from backend.utils.sanitizer import sanitize_text


class LegalPDF(FPDF):

    def __init__(self, document_type: str):
        super().__init__()
        self.document_type = document_type

    def header(self):
        self.set_font(
            "Helvetica",
            "B",
            13
        )

        self.cell(
            0,
            8,
            self.document_type.upper(),
            align="C"
        )

        self.ln(10)

    def footer(self):
        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "I",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-generated draft",
            align="C"
        )


def create_txt(content: str) -> bytes:

    return sanitize_text(
        content
    ).encode("utf-8")


def _add_docx_footer(
    document: Document
) -> None:

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            "LegalEase - AI-generated draft. "
            "Review before use."
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(8)


def create_docx(
    content: str,
    document_type: str,
    logo_path: Optional[str] = None
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Optional logo
    if logo_path:

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            paragraph.add_run().add_picture(
                logo_path,
                width=Inches(1.2)
            )

        except Exception:
            pass

    # Title
    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        document_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    # Brand
    brand = document.add_paragraph()

    brand.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = brand.add_run(
        "LegalEase"
    )

    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)

    # Document body
    for raw_line in sanitize_text(
        content
    ).splitlines():

        line = raw_line.strip()

        if not line:

            document.add_paragraph()

            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)

        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

        if (
            line.isupper()
            or line.startswith(
                tuple(
                    f"{i}."
                    for i in range(1, 100)
                )
            )
        ):
            run.bold = True

    _add_docx_footer(document)

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


def create_pdf(
    content: str,
    document_type: str
) -> bytes:

    pdf = LegalPDF(
        document_type
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.set_margins(
        18,
        20,
        18
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        size=10.5
    )

    for raw_line in sanitize_text(
        content
    ).splitlines():

        line = raw_line.strip()

        if not line:

            pdf.ln(4)

            continue

        pdf.multi_cell(
            0,
            6.5,
            line
        )

        pdf.ln(1)

    return bytes(
        pdf.output()
    )